"""Export node-test results to CSV, JSON, a colour PNG scorecard and an XLSX workbook."""
from __future__ import annotations

import csv
import datetime
import json
import os
import re
from pathlib import Path
from typing import Any

from node_groups import classify

PLATFORMS = ["ChatGPT", "Claude", "Gemini", "Netflix", "Disney+", "YouTube Premium"]
SENSITIVE_KEY = re.compile(r"password|passwd|secret|token|credential|cookie|authorization|uuid|订阅地址|订阅链接|代理密码", re.I)
URL_RE = re.compile(r"\b[a-z][a-z0-9+.-]*://\S+", re.I)
KEEP_URL_KEYS = {"test_url", "source_url", "url_source", "speed_source_url"}
TZ = datetime.timezone(datetime.timedelta(hours=8))


def _clean(value: Any, key: str = "") -> Any:
    """Drop credentials and subscription URLs while retaining report evidence."""
    if isinstance(value, dict):
        return {str(k): _clean(v, str(k)) for k, v in value.items() if not SENSITIVE_KEY.search(str(k))}
    if isinstance(value, list):
        return [_clean(v, key) for v in value]
    if isinstance(value, str) and "://" in value and key.lower() not in KEEP_URL_KEYS:
        return URL_RE.sub("[已隐藏URL]", value)
    return value


def _csv_safe(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


# --------------------------------------------------------------------------- shared judgements

def row_order(row):
    info = classify({'name': row.get('node', ''), 'type': row.get('protocol', '')})
    country = info['country'] if info['country'] != '未知' else (row.get('exit_country') or 'ZZ')
    lat = row.get('latency_ms') if isinstance(row.get('latency_ms'), (int, float)) else 1e9
    return (country or 'ZZ', float(info['multiplier']), row.get('chain', ''), lat, row.get('node', ''))


def traffic_cost(rows):
    raw = charged = 0.0
    for row in rows:
        down = row.get('traffic_download_bytes')
        if down is None:
            down = round(((row.get('download_mb') or 0) + (row.get('small_download_mb') or 0)) * 1e6)
        up = row.get('traffic_upload_bytes') or 0
        mb = (down + up + (row.get('previous_test_traffic_bytes') or 0)) / 1e6
        multiplier = row_order(row)[1]
        row['traffic_multiplier'] = multiplier
        row['traffic_mb'] = round(mb, 3)
        row['charged_traffic_mb'] = round(mb * multiplier, 3)
        raw += mb
        charged += mb * multiplier
    return {'download_mb': round(raw, 3), 'charged_mb': round(charged, 3), 'charged_gb': round(charged / 1000, 4)}


def network_quality(r):
    speed, delay, loss, jitter = r.get('speed_mbps'), r.get('latency_ms'), r.get('loss_pct'), r.get('jitter_ms')
    if speed is None:
        return ('待确认', '未完成下载测速；可单独重测') if delay is not None else ('暂不可用', '本轮连接探测失败；可重测确认')
    if speed < 10 or (loss is not None and loss >= 66):
        return ('较差', '加载慢或连接失败较多，不作首选')
    if delay is None:
        return ('待确认', '有下载速度，缺少延迟数据')
    if delay > 600:
        return ('响应偏慢', '下载可用；网页与交互响应偏慢')
    if speed < 20:
        return ('备用', '速度不足20 Mbps，轻量使用或备用')
    if (loss or 0) > 0 or (jitter or 0) > 150:
        return ('有波动', '速度达标，延迟有波动；建议重测确认')
    if speed >= 50 and delay <= 250 and jitter is not None and jitter <= 60:
        return ('优秀', '适合网页、AI交互及大文件下载')
    if delay <= 400:
        return ('良好', '适合网页和AI交互；平台可用性另看检测结果')
    return ('可用', '适合下载与视频，交互响应一般')


def speed_level(row, key='speed_mbps'):
    value = row.get(key)
    if not isinstance(value, (int, float)):
        return '未测'
    return '低速' if value < 10 else ('备用' if value < 20 else '达标')


def sampling(row, prefix='download'):
    value = row.get('speed_mbps' if prefix == 'download' else 'upload_mbps')
    if value is None:
        return '未测'
    return '完整' if row.get(f'{prefix}_complete') else '限时采样'


def _eligible(r):
    return isinstance(r.get('speed_mbps'), (int, float)) and r['speed_mbps'] >= 20 and not r.get('error') and (r.get('loss_pct') is None or r['loss_pct'] <= 34)


def country_candidates(rows):
    pools = {}
    for row in rows:
        if not _eligible(row):
            continue
        country = row_order(row)[0]
        if country in ('未知', 'ZZ', ''):
            continue
        pools.setdefault((row.get('chain', ''), country), []).append(row)
    result = []
    for _, group in sorted(pools.items()):
        group.sort(key=lambda r: (row_order(r)[1], r.get('loss_pct') if r.get('loss_pct') is not None else 100,
                                  r.get('latency_ms') if r.get('latency_ms') is not None else 1e9, -r['speed_mbps']))
        result.append(group[0])
    return result


def top_nodes(rows):
    result = []
    for chain in sorted({r.get('chain', '') for r in rows}):
        eligible = [r for r in rows if r.get('chain', '') == chain and _eligible(r)]
        eligible.sort(key=lambda r: (-r['speed_mbps'], r.get('latency_ms') or 1e9))
        result.extend(eligible[:5])
    return result


def report_stem(data):
    meta = data.get('meta') or {}
    airport = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(meta.get('profile') or '未命名机场'))
    airport = re.sub(r'\.ya?ml$', '', airport, flags=re.I).strip(' .')[:90] or '未命名机场'
    try:
        date = datetime.datetime.fromisoformat(meta.get('started')).astimezone(TZ)
    except Exception:
        date = datetime.datetime.now(TZ)
    return airport + '_' + date.strftime('%Y%m%d-%H%M%S')


def _status(row, p):
    return ((row.get('checks') or {}).get(p) or {}).get('status', '未测')


# --------------------------------------------------------------------------- table definitions

MAIN_COLUMNS = [
    ('节点', 46, lambda r: r.get('node', '')),
    ('链路', 12, lambda r: r.get('chain', '')),
    ('协议', 11, lambda r: r.get('protocol', '')),
    ('延迟(ms)', 10, lambda r: r.get('latency_ms')),
    ('抖动(ms)', 10, lambda r: r.get('jitter_ms')),
    ('延迟失败率(%)', 12, lambda r: r.get('loss_pct')),
    ('下载(Mbps)', 12, lambda r: r.get('speed_mbps')),
    ('下载平均(Mbps)', 13, lambda r: r.get('speed_avg_mbps')),
    ('下载采样', 10, lambda r: sampling(r, 'download')),
    ('上传(Mbps)', 12, lambda r: r.get('upload_mbps')),
    ('上传采样', 10, lambda r: sampling(r, 'upload')),
    *[(f'{p}', 16, (lambda p: lambda r: _status(r, p))(p)) for p in PLATFORMS],
    ('网络表现', 11, lambda r: network_quality(r)[0]),
    ('使用建议', 34, lambda r: network_quality(r)[1]),
    ('速率分级', 10, lambda r: speed_level(r)),
    ('下载测速源', 20, lambda r: r.get('speed_source', '')),
    ('上传测速源', 22, lambda r: r.get('upload_source', '')),
    ('线程数', 8, lambda r: r.get('download_streams')),
    ('下载量(MB)', 11, lambda r: r.get('download_mb')),
    ('上传量(MB)', 11, lambda r: r.get('upload_mb')),
    ('下载耗时(秒)', 11, lambda r: r.get('download_elapsed_s')),
    ('提前结束原因', 26, lambda r: r.get('download_stop_reason', '')),
    ('上传结束原因', 26, lambda r: r.get('upload_stop_reason', '')),
    ('入口IP', 16, lambda r: r.get('entry_ip', '')),
    ('入口国家', 9, lambda r: r.get('entry_country', '')),
    ('入口ASN', 10, lambda r: r.get('entry_asn', '')),
    ('出口IP', 16, lambda r: r.get('exit_ip', '')),
    ('出口国家', 9, lambda r: r.get('exit_country', '')),
    ('出口ASN', 10, lambda r: r.get('exit_asn', '')),
    ('同出口节点数', 11, lambda r: r.get('same_exit_count')),
    ('错误', 34, lambda r: r.get('error', '') or r.get('upload_error', '')),
    ('自动分类', 34, lambda r: r.get('group_label', '')),
    ('同类数量', 9, lambda r: r.get('group_size')),
    ('同类节点（未实测）', 40, lambda r: '\n'.join(r.get('group_members') or [])),
    ('测量时间', 20, lambda r: r.get('measured_at', '')),
    ('流量倍率', 9, lambda r: r.get('traffic_multiplier')),
    ('测速流量(MB)', 12, lambda r: r.get('traffic_mb')),
    ('预计扣费(MB)', 12, lambda r: r.get('charged_traffic_mb')),
]

NOTE = ('延迟为 Mihomo unified-delay（不含握手，约等于经节点访问 Google 的一个往返）；失败率是3次延迟探测的失败比例，不是ICMP丢包。'
        '下载/上传为去掉开头爬升期后的稳态速率，平均值另列；不同测速源的结果只在同一地区内可比。'
        '入口/出口信息不代表完整路由。平台“支持地区”只说明出口地区未被拦截，不代表账号可用。')


def _csv_rows(data):
    yield [c[0] for c in MAIN_COLUMNS]
    for row in data.get('rows', []):
        yield [_csv_safe(c[2](row)) for c in MAIN_COLUMNS]


def status_colors(status: str):
    s = str(status or '').lower()
    if any(k in s for k in ('风控', 'captcha', 'risk')):
        return 'FDE2E2', '8F2020'
    if '限制' in s:
        return 'FFF0CC', '815900'
    if '仅' in s:
        return 'FFF6DD', '7A5A00'
    if any(k in s for k in ('超时', 'timeout')):
        return 'E9E2F7', '523A7A'
    if any(k in s for k in ('待确认', '未测', 'pending', 'unknown')):
        return 'E8EDF2', '526477'
    if any(k in s for k in ('支持', '可访问', '可用', '正常')):
        return 'DDF1E5', '1E6840'
    return 'EEF1F4', '465564'


# --------------------------------------------------------------------------- xlsx

def _make_xlsx(data, dest: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    rows = data.get('rows', [])
    meta = data.get('meta') or {}
    font = 'SimSun'
    header_fill = PatternFill('solid', fgColor='234766')
    thin = Side(style='thin', color='E2E8ED')

    def sheet(ws, headers, widths, body, table_name, row_height=30, status_cols=()):
        ws.append(headers)
        for r in body:
            ws.append(r)
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for cell in ws[1]:
            cell.font = Font(name=font, size=10, bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 34
        for row_cells in ws.iter_rows(min_row=2, max_row=len(body) + 1):
            ws.row_dimensions[row_cells[0].row].height = row_height
            for cell in row_cells:
                cell.font = Font(name=font, size=10, color='263746')
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.border = Border(bottom=thin)
            for idx in status_cols:
                cell = row_cells[idx]
                fill, ink = status_colors(cell.value)
                cell.fill = PatternFill('solid', fgColor=fill)
                cell.font = Font(name=font, size=10, bold=True, color=ink)
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        if body:
            ref = f'A1:{get_column_letter(len(headers))}{len(body) + 1}'
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
            ws.add_table(table)
        ws.freeze_panes = 'B2'
        note_row = len(body) + 3
        ws.cell(note_row, 1, f"测试时间：{meta.get('started', '')}；模式：{meta.get('mode', '')}。" + NOTE)
        ws.cell(note_row, 1).font = Font(name=font, size=9, italic=True, color='526477')
        ws.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical='center')
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=min(10, len(headers)))
        ws.row_dimensions[note_row].height = 60

    wb = Workbook()
    picks = wb.active
    picks.title = '推荐节点'
    best = []
    for r in top_nodes(rows):
        best.append(['达标速度前5', r.get('chain', ''), r['node'], row_order(r)[0], row_order(r)[1], r['speed_mbps'], r.get('upload_mbps'), r.get('latency_ms'), sampling(r)])
    for r in country_candidates(rows):
        best.append(['各国低倍率', r.get('chain', ''), r['node'], row_order(r)[0], row_order(r)[1], r['speed_mbps'], r.get('upload_mbps'), r.get('latency_ms'), sampling(r)])
    if best:
        sheet(picks, ['推荐类型', '链路', '节点', '国家', '倍率', '下载(Mbps)', '上传(Mbps)', '延迟(ms)', '下载采样'],
              [16, 12, 50, 8, 8, 12, 12, 10, 10], best, 'TopSpeeds')
    else:
        picks['A1'] = '暂无达到20 Mbps的有效节点推荐'

    main = wb.create_sheet('测评结果')
    status_start = [c[0] for c in MAIN_COLUMNS].index(PLATFORMS[0])
    sheet(main, [c[0] for c in MAIN_COLUMNS], [c[1] for c in MAIN_COLUMNS], [[c[2](r) for c in MAIN_COLUMNS] for r in rows],
          'NodeResults', 36, range(status_start, status_start + len(PLATFORMS)))

    detail = wb.create_sheet('检测详情')
    drows = []
    for r in rows:
        for p in PLATFORMS:
            c = (r.get('checks') or {}).get(p) or {}
            drows.append([r.get('node', ''), r.get('chain', ''), p, c.get('status', '未测'), c.get('region', ''), c.get('detail', ''), c.get('test_url', ''), c.get('source_url', '')])
    sheet(detail, ['节点', '链路', '平台', '状态', '地区', '证据/详情', '测试URL', '判据来源'], [46, 12, 16, 16, 8, 70, 40, 40],
          drows, 'CheckDetails', 48, (3,))
    wb.save(dest)


# --------------------------------------------------------------------------- png

def _font(size):
    from PIL import ImageFont
    candidates = [r'C:\Windows\Fonts\simsun.ttc', r'C:\Windows\Fonts\msyh.ttc', r'C:\Windows\Fonts\simhei.ttf',
                  '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc',
                  '/System/Library/Fonts/PingFang.ttc']
    for path in candidates:
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size, index=0)
            except OSError:
                continue
    return ImageFont.load_default()


def _make_png(data: dict[str, Any], dest: Path) -> None:
    from PIL import Image, ImageDraw

    body, small, head, title = _font(20), _font(17), _font(21), _font(32)
    rows = sorted(data.get("rows", []), key=row_order)
    meta = data.get("meta") or {}
    labels = ["节点", "链路", "网络表现", "使用建议", "延迟\nms", "下载\nMbps", "上传\nMbps", *PLATFORMS]
    widths = [330, 125, 150, 290, 100, 120, 120, 175, 175, 175, 175, 175, 200]
    width = sum(widths) + 80
    margin, top = 40, 42
    header_h, row_h = 65, 110
    costs = traffic_cost(rows)
    note_lines = [f"时间：{meta.get('started') or '未提供'}    模式：{meta.get('mode') or '未提供'}    线程：{meta.get('streams') or 1}",
                  f"测速流量合计 {costs['download_mb']:.2f} MB（含上传）；按倍率预计扣费 {costs['charged_mb']:.2f} MB（{costs['charged_gb']:.4f} GB）。",
                  "延迟=Mihomo统一延迟；速度=去掉爬升期后的稳态速率；空白表示未测；平台“支持地区”不代表账号可用。",
                  '推荐门槛：下载至少20 Mbps；低于10 Mbps为低速，10–20 Mbps为备用。']
    tops = top_nodes(rows)
    if not tops:
        note_lines.append('暂无达到20 Mbps的有效节点推荐。')
    for chain in sorted({r.get('chain', '') for r in tops}):
        note_lines.append(chain + ' · 下载速度前5')
        for i, r in enumerate([r for r in tops if r.get('chain', '') == chain], 1):
            up = f" / 上传 {r['upload_mbps']:.2f}" if isinstance(r.get('upload_mbps'), (int, float)) else ''
            note_lines.append(f"{i}. {r['node']}    下载 {r['speed_mbps']:.2f}{up} Mbps" + ('（限时采样）' if not r.get('download_complete') else ''))
    cands = country_candidates(rows)
    if cands:
        note_lines.append('各国低倍率候选（均达到20 Mbps）')
        for r in cands:
            country, multiplier, *_ = row_order(r)
            note_lines.append(f"{country} · {r.get('chain', '')} · {multiplier:g}x · {r['speed_mbps']:.2f} Mbps · {r['node']}")
    height = top + 110 + 32 * len(note_lines) + header_h + max(1, len(rows)) * row_h + 30
    im = Image.new("RGB", (width, height), "#F5F7FA")
    draw = ImageDraw.Draw(im)
    draw.text((margin, top), str(meta.get("profile") or "未命名机场") + " · 节点测评", font=title, fill="#17324D")
    y = top + 56
    for line in note_lines:
        draw.text((margin, y), re.sub(r"[\U0001F1E6-\U0001F1FF]", "", line), font=small, fill="#526477")
        y += 32
    y += 8
    x = margin
    for label, colw in zip(labels, widths):
        draw.rounded_rectangle((x, y, x + colw - 3, y + header_h), radius=7, fill="#234766")
        draw.multiline_text((x + 9, y + 10), label, font=head, fill="white", spacing=2)
        x += colw
    y += header_h

    def wrap(text, font, max_px, max_lines=3):
        text = re.sub(r"[\U0001F1E6-\U0001F1FF]", "", str(text or "未测"))
        lines, cur = [], ""
        for ch in text:
            if ch == '\n':
                lines.append(cur)
                cur = ''
                continue
            trial = cur + ch
            if draw.textbbox((0, 0), trial, font=font)[2] <= max_px or not cur:
                cur = trial
            else:
                lines.append(cur)
                cur = ch
        if cur:
            lines.append(cur)
        if len(lines) > max_lines:
            lines = lines[:max_lines]
            lines[-1] = lines[-1][:-1] + "…"
        return lines

    for i, row in enumerate(rows):
        x = margin
        bg = "#FFFFFF" if i % 2 == 0 else "#F0F4F7"
        name = row.get("node", "")
        if row.get('group_size', 1) > 1:
            name += f"（同类{row['group_size']}个，仅此节点实测）"
        vals = [name, row.get("chain", ""), *network_quality(row), row.get("latency_ms"), row.get("speed_mbps"), row.get("upload_mbps")]
        vals += [_status(row, p) for p in PLATFORMS]
        for j, (val, colw) in enumerate(zip(vals, widths)):
            if j >= 7:
                fill, ink = ('#' + c for c in status_colors(str(val)))
            else:
                fill, ink = bg, "#263746"
            draw.rounded_rectangle((x, y + 2, x + colw - 3, y + row_h - 2), radius=5, fill=fill)
            if j in (4, 5, 6):
                val = "未测" if val is None else f"{val:g}"
                if j == 5 and row.get('speed_mbps') is not None:
                    val += "\n" + speed_level(row) + ("" if row.get('download_complete') else "\n限时")
                if j == 6 and row.get('upload_mbps') is not None and not row.get('upload_complete'):
                    val += "\n限时"
            lines = wrap(str(val or "未测"), body, colw - 20, 4)
            yy = y + max(8, (row_h - len(lines) * 24) // 2)
            for line in lines:
                draw.text((x + 9, yy), line, font=body, fill=ink)
                yy += 24
            x += colw
        y += row_h
    im.save(dest, format="PNG", optimize=True)


# --------------------------------------------------------------------------- entry point

def export_report(data: dict[str, Any], outdir) -> dict[str, str]:
    """Write a sanitized report bundle. Missing optional libraries only skip that format."""
    output = Path(outdir).resolve()
    output.mkdir(parents=True, exist_ok=True)
    safe = _clean(data)
    safe['rows'] = sorted(safe.get('rows', []), key=row_order)
    for row in safe['rows']:
        row['speed_level'] = speed_level(row)
        row['network_quality'], row['usage_advice'] = network_quality(row)
    safe.setdefault('meta', {})['traffic_cost'] = traffic_cost(safe['rows'])
    stem = report_stem(safe)
    files = {}
    json_path = output / (stem + "_原始结果.json")
    json_path.write_text(json.dumps(safe, ensure_ascii=False, indent=2), encoding="utf-8")
    files['json'] = json_path
    csv_path = output / (stem + "_节点测评.csv")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(_csv_rows(safe))
    files['csv'] = csv_path
    errors = []
    for key, maker, suffix in (('xlsx', _make_xlsx, '_节点测评.xlsx'), ('png', _make_png, '_节点测评.png')):
        path = output / (stem + suffix)
        try:
            maker(safe, path)
            files[key] = path
        except Exception as exc:  # noqa: BLE001 - other formats are still useful
            errors.append(f'{key}: {type(exc).__name__}: {exc}')
    if errors:
        (output / '导出错误.txt').write_text('\n'.join(errors), encoding='utf-8')
    order = ('xlsx', 'png', 'csv', 'json')
    manifest = {k: files[k].name for k in order if k in files}
    (output / "报告文件.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    return {k: str(files[k]) for k in order if k in files}
